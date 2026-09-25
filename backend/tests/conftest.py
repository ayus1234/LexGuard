import io
import sys
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

# Ensure backend root and app directories are in sys.path
backend_dir = Path(__file__).resolve().parent.parent
project_dir = backend_dir.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))
if str(project_dir) not in sys.path:
    sys.path.insert(0, str(project_dir))

try:
    from backend.app.main import app
    from backend.app.core.config import settings
except ImportError:
    from app.main import app
    from app.core.config import settings


@pytest.fixture(scope="session")
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Generate a multi-page valid PDF with legal clauses and sections."""
    import pymupdf
    doc = pymupdf.open()

    # Page 1
    p1 = doc.new_page()
    p1.insert_text(
        (50, 72),
        "ARTICLE 1: MUTUAL NON-DISCLOSURE AGREEMENT\n\n"
        "This Mutual Non-Disclosure Agreement is executed between LexGuard Corp and Client.\n"
        "All confidential proprietary information disclosed under this agreement shall remain protected.",
        fontsize=11,
    )

    # Page 2
    p2 = doc.new_page()
    p2.insert_text(
        (50, 72),
        "§ 8.3 Limitation of Liability\n\n"
        "In no event shall either party's aggregate liability exceed the total amounts paid hereunder.\n"
        "Neither party shall be liable for consequential, incidental, or special damages.",
        fontsize=11,
    )

    pdf_buffer = io.BytesIO()
    doc.save(pdf_buffer)
    doc.close()
    return pdf_buffer.getvalue()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generate a valid DOCX with headings and table data."""
    import docx
    doc = docx.Document()
    doc.add_heading("ARTICLE 4: INTELLECTUAL PROPERTY", level=1)
    doc.add_paragraph(
        "All work product, discoveries, patents, and copyrightable materials developed under this Statement of Work "
        "shall be deemed work made for hire owned exclusively by the Company."
    )
    doc.add_heading("SECTION 9.1: FEES AND PAYMENT", level=2)
    doc.add_paragraph("Customer agrees to remit payment within 30 days of receipt of invoice.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Phase 1 Completion"
    table.cell(1, 1).text = "$25,000 USD"

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    return docx_buffer.getvalue()


@pytest.fixture
def sample_txt_bytes() -> bytes:
    """Generate a valid plain text legal document."""
    content = (
        "MASTER SERVICES AGREEMENT\n\n"
        "ARTICLE 2: SCOPE OF SERVICES\n"
        "Provider agrees to perform the technical advisory and software development services.\n\n"
        "§ 5.4 Indemnification Obligations\n"
        "Provider shall defend and indemnify Customer against any third-party claims alleging infringement.\n"
    )
    return content.encode("utf-8")


@pytest.fixture
def masqueraded_file_bytes() -> bytes:
    """Fake PDF containing PNG magic bytes."""
    return b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"


@pytest.fixture
def corrupted_pdf_bytes() -> bytes:
    """File that starts with %PDF- but is truncated / contains unparseable binary junk."""
    return b"%PDF-1.7\nCorrupted binary junk \x00\xff\xfe\x01\x02%%EOF"
