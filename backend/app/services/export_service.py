"""
LexGuard Document Export Service.
Generates professional PDF and DOCX export dossiers for the Lawyer Preparation Brief.
Strictly adheres to zero-retention privacy, non-legal advice disclaimers, and verbatim grounding citations.
"""

import io
from datetime import datetime, timezone
from typing import TYPE_CHECKING, List, Optional

if TYPE_CHECKING:
    from app.schemas.brief import LawyerBrief
else:
    try:
        from backend.app.schemas.brief import LawyerBrief
    except ImportError:
        from app.schemas.brief import LawyerBrief


class ExportService:
    """
    Renders structured Lawyer Preparation Briefs into high-fidelity PDF and DOCX documents.
    """

    @staticmethod
    def generate_brief_docx(brief: LawyerBrief) -> bytes:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = docx.Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Palette
        NAVY = RGBColor(15, 23, 42)      # #0F172A
        BLUE = RGBColor(2, 132, 199)     # #0284C7
        SLATE = RGBColor(100, 116, 139)  # #64748B
        DARK = RGBColor(51, 65, 85)      # #334155

        # Document Header
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("LEXGUARD — COUNSEL PREPARATION BRIEF")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = NAVY

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(
            "Confidential Legal Work-Product Preparation Dossier • Educational Assistance Only"
        )
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(9)
        sub_run.font.italic = True
        sub_run.font.color.rgb = SLATE

        # Metadata Table
        meta_table = doc.add_table(rows=5, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        rows_data = [
            ("Target Document:", brief.document_title),
            ("Document Type:", brief.document_type),
            ("Governing Jurisdiction:", brief.jurisdiction or "Not explicitly stated in document"),
            ("Generated Timestamp:", brief.generated_at.strftime("%B %d, %Y at %H:%M:%S UTC")),
            ("Dossier Reference ID:", f"LG-{brief.document_id[:16].upper()}"),
        ]

        for i, (label, val) in enumerate(rows_data):
            cell_lbl = meta_table.cell(i, 0)
            cell_val = meta_table.cell(i, 1)
            cell_lbl.width = Inches(2.0)
            cell_val.width = Inches(5.0)

            p_l = cell_lbl.paragraphs[0]
            r_l = p_l.add_run(label)
            r_l.font.name = "Arial"
            r_l.font.size = Pt(9.5)
            r_l.font.bold = True
            r_l.font.color.rgb = NAVY

            p_v = cell_val.paragraphs[0]
            r_v = p_v.add_run(val)
            r_v.font.name = "Arial"
            r_v.font.size = Pt(9.5)
            r_v.font.color.rgb = DARK

        doc.add_paragraph()  # Spacing

        # Section 1: Executive Summary
        h1 = doc.add_heading("1. Executive Document Summary", level=1)
        h1.runs[0].font.color.rgb = NAVY
        p_exec = doc.add_paragraph(brief.executive_summary)
        p_exec.runs[0].font.name = "Arial"
        p_exec.runs[0].font.size = Pt(10)
        p_exec.runs[0].font.color.rgb = DARK

        # Section 2: Key Information Table
        if brief.key_information:
            doc.add_paragraph()
            h2 = doc.add_heading("2. Key Commercial & Legal Information", level=1)
            h2.runs[0].font.color.rgb = NAVY

            k_table = doc.add_table(rows=1, cols=4)
            k_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = k_table.rows[0].cells
            headers = ["Category", "Descriptor", "Document-Stated Value", "Source Anchor"]
            for idx, name in enumerate(headers):
                hdr_cells[idx].paragraphs[0].add_run(name).font.bold = True

            for item in brief.key_information:
                row_cells = k_table.add_row().cells
                row_cells[0].paragraphs[0].add_run(item.category)
                row_cells[1].paragraphs[0].add_run(item.label)
                row_cells[2].paragraphs[0].add_run(item.value or "Not stated")
                ref_txt = item.source_reference or (f"Page {item.page}" if item.page else "—")
                row_cells[3].paragraphs[0].add_run(ref_txt)

        # Section 3: Potential Areas for Attention
        if brief.attention_areas:
            doc.add_paragraph()
            h3 = doc.add_heading("3. Potential Areas for Attention (Counsel Review)", level=1)
            h3.runs[0].font.color.rgb = NAVY

            for item in brief.attention_areas:
                p_item = doc.add_paragraph()
                r_badge = p_item.add_run(f"[{item.review_level.replace('_', ' ').title()}] ")
                r_badge.font.bold = True
                r_badge.font.color.rgb = BLUE
                r_title = p_item.add_run(item.title)
                r_title.font.bold = True
                r_title.font.color.rgb = NAVY

                p_desc = doc.add_paragraph()
                p_desc.add_run(f"Contract Provision: {item.description}\n").font.size = Pt(9.5)
                p_desc.add_run(f"Why Attention is Warranted: {item.why_it_matters}\n").font.size = Pt(9.5)
                if item.source_reference or item.page:
                    ref_line = f"Source Reference: {item.source_reference or ''} (Page {item.page or '—'})"
                    r_ref = p_desc.add_run(ref_line)
                    r_ref.font.italic = True
                    r_ref.font.size = Pt(9)
                    r_ref.font.color.rgb = SLATE

        # Section 4: Negotiation Points
        if brief.negotiation_points:
            doc.add_paragraph()
            h4 = doc.add_heading("4. Potential Negotiation Discussion Points", level=1)
            h4.runs[0].font.color.rgb = NAVY

            for neg in brief.negotiation_points:
                p_neg = doc.add_paragraph()
                r_nt = p_neg.add_run(f"Topic: {neg.title}")
                r_nt.font.bold = True
                r_nt.font.color.rgb = NAVY

                p_body = doc.add_paragraph()
                p_body.add_run(f"Current Provision: {neg.current_provision}\n").font.size = Pt(9.5)
                p_body.add_run(f"Discussion Point: {neg.discussion_point}\n").font.size = Pt(9.5)
                if neg.suggested_compromise:
                    p_body.add_run(f"Suggested Compromise Formulation: \"{neg.suggested_compromise}\"\n").font.bold = True
                if neg.market_baseline:
                    p_body.add_run(f"Market Baseline: {neg.market_baseline}\n").font.size = Pt(9)

        # Section 5: Prioritized Questions for Counsel
        if brief.counsel_questions:
            doc.add_paragraph()
            h5 = doc.add_heading("5. Prioritized Questions for Counsel Consultation", level=1)
            h5.runs[0].font.color.rgb = NAVY

            for q in brief.counsel_questions:
                p_q = doc.add_paragraph()
                r_pri = p_q.add_run(f"[{q.priority} Priority] ")
                r_pri.font.bold = True
                r_pri.font.color.rgb = BLUE
                r_topic = p_q.add_run(q.agenda_topic)
                r_topic.font.bold = True
                r_topic.font.color.rgb = NAVY

                p_qbody = doc.add_paragraph()
                p_qbody.add_run(f"Context: {q.why_discuss}\n").font.size = Pt(9.5)
                if q.suggested_phrasing:
                    p_qbody.add_run(f"Suggested Question Phrasing: \"{q.suggested_phrasing}\"\n").font.italic = True
                if q.contract_citation:
                    p_qbody.add_run(f"Contract Anchor: {q.contract_citation} (Page {q.page or '—'})\n").font.size = Pt(9)

        # Section 6: Action Checklist
        if brief.checklist:
            doc.add_paragraph()
            h6 = doc.add_heading("6. Action Checklist (\"Before You Proceed\")", level=1)
            h6.runs[0].font.color.rgb = NAVY

            for chk in brief.checklist:
                p_chk = doc.add_paragraph()
                box = "[X] " if chk.completed else "[ ] "
                p_chk.add_run(box).font.bold = True
                r_chk_title = p_chk.add_run(f"{chk.title} ")
                r_chk_title.font.bold = True
                r_chk_title.font.size = Pt(9.5)
                p_chk.add_run(f"({chk.badge_text})\n").font.color.rgb = BLUE
                p_chk.add_run(f"    Directive Anchor: {chk.citation}").font.size = Pt(9)

        # Section 7: Source References & Verbatim Citations
        if brief.citations:
            doc.add_paragraph()
            h7 = doc.add_heading("7. Verbatim Source Citations Ledger", level=1)
            h7.runs[0].font.color.rgb = NAVY

            for c in brief.citations:
                p_cit = doc.add_paragraph()
                status = "[VERIFIED]" if c.verified else "[UNVERIFIED]"
                p_cit.add_run(f"{status} Page {c.page or '—'} • Section {c.section or '—'}\n").font.bold = True
                r_quote = p_cit.add_run(f"\"{c.quoted_text}\"")
                r_quote.font.italic = True
                r_quote.font.size = Pt(9)
                r_quote.font.color.rgb = SLATE

        # Section 8: Disclaimer
        doc.add_paragraph()
        doc.add_heading("Mandatory Legal Disclaimer", level=2).runs[0].font.color.rgb = NAVY
        p_disc = doc.add_paragraph()
        r_disc = p_disc.add_run(brief.disclaimer)
        r_disc.font.name = "Arial"
        r_disc.font.size = Pt(8.5)
        r_disc.font.italic = True
        r_disc.font.color.rgb = SLATE

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def generate_brief_pdf(brief: LawyerBrief) -> bytes:
        import pymupdf

        doc = pymupdf.open()
        PAGE_WIDTH = 595.32   # A4 Width in points
        PAGE_HEIGHT = 841.92  # A4 Height in points
        MARGIN_LEFT = 45.0
        MARGIN_RIGHT = PAGE_WIDTH - 45.0
        MARGIN_TOP = 45.0
        MARGIN_BOTTOM = PAGE_HEIGHT - 45.0
        CONTENT_WIDTH = MARGIN_RIGHT - MARGIN_LEFT

        current_page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        y = MARGIN_TOP

        def check_page_break(needed_height: float) -> None:
            nonlocal current_page, y
            if y + needed_height > MARGIN_BOTTOM:
                current_page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
                y = MARGIN_TOP

        # Header Banner on Page 1
        banner_height = 42.0
        current_page.draw_rect(
            pymupdf.Rect(MARGIN_LEFT, y, MARGIN_RIGHT, y + banner_height),
            color=(0.06, 0.09, 0.16),
            fill=(0.06, 0.09, 0.16),
        )
        current_page.insert_text(
            (MARGIN_LEFT + 12, y + 26),
            "LEXGUARD — COUNSEL PREPARATION BRIEF",
            fontsize=13,
            color=(1, 1, 1),
            fontname="helv",
        )
        y += banner_height + 12

        # Subhead & Reference Bar
        sub_text = (
            f"Target: {brief.document_title} • Jurisdiction: {brief.jurisdiction or 'Unspecified'} • "
            f"Ref: LG-{brief.document_id[:12].upper()} • Generated: {brief.generated_at.strftime('%Y-%m-%d')}"
        )
        current_page.insert_text(
            (MARGIN_LEFT, y),
            sub_text,
            fontsize=8,
            color=(0.39, 0.45, 0.55),
            fontname="helv",
        )
        y += 16
        current_page.draw_line(
            pymupdf.Point(MARGIN_LEFT, y),
            pymupdf.Point(MARGIN_RIGHT, y),
            color=(0.8, 0.84, 0.88),
            width=0.75,
        )
        y += 16

        def draw_section_header(title: str) -> None:
            nonlocal y
            check_page_break(32)
            current_page.insert_text(
                (MARGIN_LEFT, y + 10),
                title.upper(),
                fontsize=10,
                color=(0.01, 0.52, 0.78),
                fontname="helv",
            )
            y += 16
            current_page.draw_line(
                pymupdf.Point(MARGIN_LEFT, y),
                pymupdf.Point(MARGIN_RIGHT, y),
                color=(0.8, 0.84, 0.88),
                width=0.5,
            )
            y += 12

        def draw_wrapped_paragraph(text: str, fontsize: float = 9.0, is_bold: bool = False, color=(0.1, 0.15, 0.22)) -> None:
            nonlocal y
            # Approximate height estimation based on characters per line
            chars_per_line = int(CONTENT_WIDTH / (fontsize * 0.52))
            lines = []
            for paragraph in text.split("\n"):
                words = paragraph.split(" ")
                curr_line = ""
                for w in words:
                    if len(curr_line) + len(w) + 1 <= chars_per_line:
                        curr_line = f"{curr_line} {w}".strip()
                    else:
                        lines.append(curr_line)
                        curr_line = w
                if curr_line:
                    lines.append(curr_line)

            line_height = fontsize * 1.35
            needed = len(lines) * line_height + 6
            check_page_break(needed)

            fname = "hebo" if is_bold else "helv"
            for line in lines:
                current_page.insert_text((MARGIN_LEFT, y), line, fontsize=fontsize, color=color, fontname=fname)
                y += line_height
            y += 4

        # 1. Executive Summary
        draw_section_header("1. Executive Document Summary")
        draw_wrapped_paragraph(brief.executive_summary, fontsize=9.0)

        # 2. Key Commercial Information
        if brief.key_information:
            draw_section_header("2. Key Commercial & Legal Information")
            for item in brief.key_information[:8]:
                val = item.value or "Not stated in agreement"
                ref = item.source_reference or (f"Page {item.page}" if item.page else "")
                line = f"• [{item.category}] {item.label}: {val} ({ref})"
                draw_wrapped_paragraph(line, fontsize=8.5)

        # 3. Attention Areas
        if brief.attention_areas:
            draw_section_header("3. Potential Areas for Attention (Counsel Review)")
            for item in brief.attention_areas:
                tag = item.review_level.replace("_", " ").title()
                header_line = f"[{tag}] {item.title}"
                draw_wrapped_paragraph(header_line, fontsize=9.0, is_bold=True, color=(0.06, 0.09, 0.16))
                desc_line = f"Provision: {item.description}\nWhy it matters: {item.why_it_matters}"
                draw_wrapped_paragraph(desc_line, fontsize=8.5)

        # 4. Negotiation Discussion Points
        if brief.negotiation_points:
            draw_section_header("4. Potential Negotiation Discussion Points")
            for neg in brief.negotiation_points:
                draw_wrapped_paragraph(f"Topic: {neg.title}", fontsize=9.0, is_bold=True)
                pts = f"Current provision: {neg.current_provision}\nDiscussion point: {neg.discussion_point}"
                if neg.suggested_compromise:
                    pts += f"\nSuggested compromise: \"{neg.suggested_compromise}\""
                draw_wrapped_paragraph(pts, fontsize=8.5)

        # 5. Questions for Counsel
        if brief.counsel_questions:
            draw_section_header("5. Prioritized Questions for Counsel Consultation")
            for q in brief.counsel_questions:
                q_head = f"[{q.priority} Priority] {q.agenda_topic}"
                draw_wrapped_paragraph(q_head, fontsize=9.0, is_bold=True)
                q_body = f"Context: {q.why_discuss}"
                if q.suggested_phrasing:
                    q_body += f"\nSuggested Question: \"{q.suggested_phrasing}\""
                if q.contract_citation:
                    q_body += f" (Anchor: {q.contract_citation})"
                draw_wrapped_paragraph(q_body, fontsize=8.5)

        # 6. Action Checklist
        if brief.checklist:
            draw_section_header("6. Action Checklist (\"Before You Proceed\")")
            for chk in brief.checklist:
                box = "[X] " if chk.completed else "[ ] "
                line = f"{box}{chk.title} ({chk.badge_text}) — Anchor: {chk.citation}"
                draw_wrapped_paragraph(line, fontsize=8.5)

        # 7. Citations Ledger
        if brief.citations:
            draw_section_header("7. Verbatim Source Citations Ledger")
            for c in brief.citations:
                status = "[VERIFIED]" if c.verified else "[UNVERIFIED]"
                cit_line = f"{status} Page {c.page or '—'} • Section {c.section or '—'}: \"{c.quoted_text}\""
                draw_wrapped_paragraph(cit_line, fontsize=8.0, color=(0.35, 0.4, 0.48))

        # Disclaimer Box
        check_page_break(60)
        y += 8
        current_page.draw_rect(
            pymupdf.Rect(MARGIN_LEFT, y, MARGIN_RIGHT, y + 48),
            color=(0.75, 0.85, 0.98),
            fill=(0.94, 0.96, 1.0),
        )
        current_page.insert_text(
            (MARGIN_LEFT + 10, y + 14),
            "MANDATORY EDUCATIONAL & NON-LEGAL ADVICE DISCLAIMER",
            fontsize=8,
            color=(0.01, 0.52, 0.78),
            fontname="hebo",
        )
        current_page.insert_textbox(
            pymupdf.Rect(MARGIN_LEFT + 10, y + 18, MARGIN_RIGHT - 10, y + 44),
            brief.disclaimer,
            fontsize=7.5,
            color=(0.2, 0.25, 0.32),
            fontname="helv",
        )

        # Add Running Footers to all pages
        total_pages = doc.page_count
        for i in range(total_pages):
            page = doc[i]
            page.draw_line(
                pymupdf.Point(MARGIN_LEFT, PAGE_HEIGHT - 35),
                pymupdf.Point(MARGIN_RIGHT, PAGE_HEIGHT - 35),
                color=(0.85, 0.88, 0.92),
                width=0.5,
            )
            page.insert_text(
                (MARGIN_LEFT, PAGE_HEIGHT - 22),
                "LexGuard Legal Intelligence • Educational Work-Product Preparation • Not Legal Advice",
                fontsize=7.5,
                color=(0.45, 0.5, 0.58),
                fontname="helv",
            )
            page_str = f"Page {i + 1} of {total_pages}"
            page.insert_text(
                (MARGIN_RIGHT - 50, PAGE_HEIGHT - 22),
                page_str,
                fontsize=7.5,
                color=(0.45, 0.5, 0.58),
                fontname="helv",
            )

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes
